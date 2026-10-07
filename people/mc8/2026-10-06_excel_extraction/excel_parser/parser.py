"""Extract an Excel workbook into the canonical JSON representation.

The parsing path intentionally uses only openpyxl. Workbooks are loaded twice:
once with formulas and once with ``data_only=True`` for cached formula results.
No formula is evaluated here.
"""

from __future__ import annotations

import hashlib
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries


SCHEMA_VERSION = "1.0.0"
PARSER_NAME = "mc8-openpyxl-canonical"


def json_value(value: Any) -> Any:
    """Return a JSON-safe value without applying display formatting."""

    if value is None or isinstance(value, (str, bool, int, float)):
        return value
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def value_type(cell: Any) -> str:
    """Map openpyxl's compact cell types to stable schema labels."""

    if cell.value is None:
        return "null"
    if cell.data_type == "f":
        return "formula"
    if cell.data_type == "e":
        return "error"
    if cell.is_date:
        return "datetime"
    if cell.data_type == "b" or isinstance(cell.value, bool):
        return "boolean"
    if cell.data_type == "n" or isinstance(cell.value, (int, float, Decimal)):
        return "number"
    return "string"


def _merged_lookup(worksheet: Any) -> dict[str, dict[str, str]]:
    lookup: dict[str, dict[str, str]] = {}
    for merged_range in worksheet.merged_cells.ranges:
        range_ref = str(merged_range)
        anchor = worksheet.cell(merged_range.min_row, merged_range.min_col).coordinate
        for row in worksheet.iter_rows(
            min_row=merged_range.min_row,
            max_row=merged_range.max_row,
            min_col=merged_range.min_col,
            max_col=merged_range.max_col,
        ):
            for cell in row:
                lookup[cell.coordinate] = {
                    "role": "anchor" if cell.coordinate == anchor else "child",
                    "range": range_ref,
                    "anchor": anchor,
                }
    return lookup


def _column_hidden(worksheet: Any, column_index: int) -> bool:
    """Respect ordinary and grouped column dimensions."""

    letter = get_column_letter(column_index)
    direct = worksheet.column_dimensions.get(letter)
    if direct is not None and bool(direct.hidden):
        return True
    for dimension in worksheet.column_dimensions.values():
        minimum = dimension.min or 0
        maximum = dimension.max or minimum
        if minimum <= column_index <= maximum and bool(dimension.hidden):
            return True
    return False


def _used_bounds(worksheet: Any) -> tuple[int, int, int, int] | None:
    """Find the smallest rectangle containing values, styles, or merged cells."""

    coordinates: list[tuple[int, int]] = []
    for row in worksheet.iter_rows(
        min_row=1,
        max_row=max(worksheet.max_row, 1),
        min_col=1,
        max_col=max(worksheet.max_column, 1),
    ):
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            if cell.value is not None or cell.has_style or cell.comment or cell.hyperlink:
                coordinates.append((cell.row, cell.column))
    for merged_range in worksheet.merged_cells.ranges:
        coordinates.extend(
            [
                (merged_range.min_row, merged_range.min_col),
                (merged_range.max_row, merged_range.max_col),
            ]
        )
    if not coordinates:
        return None
    rows = [row for row, _ in coordinates]
    columns = [column for _, column in coordinates]
    return min(columns), min(rows), max(columns), max(rows)


def _row_values(worksheet: Any, row: int, min_col: int, max_col: int) -> list[Any]:
    return [worksheet.cell(row, column).value for column in range(min_col, max_col + 1)]


def _nonempty(value: Any) -> bool:
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def _row_text(worksheet: Any, row: int, min_col: int, max_col: int) -> str:
    return " ".join(
        str(value).strip()
        for value in _row_values(worksheet, row, min_col, max_col)
        if _nonempty(value)
    )


def _has_spanning_merge(worksheet: Any, row: int, min_col: int, max_col: int) -> bool:
    return any(
        merged.min_row <= row <= merged.max_row
        and merged.min_col <= min_col
        and merged.max_col >= max_col
        for merged in worksheet.merged_cells.ranges
    )


def detect_tables(worksheet: Any, bounds: tuple[int, int, int, int] | None) -> list[dict[str, Any]]:
    """Detect classic financial-statement tables with transparent v1 rules.

    A core table is a run of rows containing at least two non-empty cells.
    Core rows separated by at most one intervening row remain in one table.
    Nearby one-cell/merged rows become titles, while rows beginning with
    ``note``, ``source``, or ``*`` become footnotes.
    """

    if bounds is None:
        return []
    min_col, min_row, max_col, max_row = bounds
    counts = {
        row: sum(_nonempty(value) for value in _row_values(worksheet, row, min_col, max_col))
        for row in range(min_row, max_row + 1)
    }
    dense_rows = [row for row, count in counts.items() if count >= 2]
    if not dense_rows:
        return []

    groups: list[list[int]] = [[dense_rows[0]]]
    for row in dense_rows[1:]:
        if row - groups[-1][-1] <= 2:
            groups[-1].append(row)
        else:
            groups.append([row])

    tables: list[dict[str, Any]] = []
    for index, group in enumerate(groups, start=1):
        start_row, end_row = group[0], group[-1]
        occupied_columns = [
            column
            for row in group
            for column in range(min_col, max_col + 1)
            if _nonempty(worksheet.cell(row, column).value)
        ]
        table_min_col = min(occupied_columns)
        table_max_col = max(occupied_columns)

        title_rows: list[int] = []
        for row in range(max(min_row, start_row - 3), start_row):
            text = _row_text(worksheet, row, table_min_col, table_max_col)
            if not text:
                continue
            lowered = text.casefold()
            if lowered.startswith(("note", "source", "*")):
                continue
            if counts[row] == 1 or _has_spanning_merge(
                worksheet, row, table_min_col, table_max_col
            ):
                title_rows.append(row)

        footnote_rows: list[int] = []
        for row in range(end_row + 1, min(max_row, end_row + 3) + 1):
            text = _row_text(worksheet, row, table_min_col, table_max_col)
            if text.casefold().startswith(("note", "source", "*")):
                footnote_rows.append(row)

        header_rows = [start_row]
        body_rows = list(range(start_row + 1, end_row + 1))
        range_ref = (
            f"{get_column_letter(table_min_col)}{start_row}:"
            f"{get_column_letter(table_max_col)}{end_row}"
        )
        title = " ".join(
            _row_text(worksheet, row, table_min_col, table_max_col) for row in title_rows
        ).strip()
        tables.append(
            {
                "id": f"table-{index}",
                "range": range_ref,
                "title": title or None,
                "row_roles": {
                    "title": title_rows,
                    "header": header_rows,
                    "body": body_rows,
                    "footnote": footnote_rows,
                },
                "cell_coordinates": [
                    worksheet.cell(row, column).coordinate
                    for row in range(start_row, end_row + 1)
                    for column in range(table_min_col, table_max_col + 1)
                ],
            }
        )
    return tables


def _extract_sheet(formula_sheet: Any, cached_sheet: Any) -> dict[str, Any]:
    bounds = _used_bounds(formula_sheet)
    merge_lookup = _merged_lookup(formula_sheet)
    cells: list[dict[str, Any]] = []
    if bounds is not None:
        min_col, min_row, max_col, max_row = bounds
        for row in range(min_row, max_row + 1):
            for column in range(min_col, max_col + 1):
                cell = formula_sheet.cell(row, column)
                cached_cell = cached_sheet.cell(row, column)
                is_formula = cell.data_type == "f"
                cached = json_value(cached_cell.value) if is_formula else None
                cells.append(
                    {
                        "coordinate": cell.coordinate,
                        "value": json_value(cell.value),
                        "value_type": value_type(cell),
                        "formula": str(cell.value) if is_formula else None,
                        "cached_value": cached,
                        "cached_value_status": (
                            "present" if is_formula and cached_cell.value is not None
                            else "missing" if is_formula
                            else "not_formula"
                        ),
                        "number_format": cell.number_format,
                        "merged": merge_lookup.get(cell.coordinate),
                        "row_hidden": bool(formula_sheet.row_dimensions[row].hidden),
                        "column_hidden": _column_hidden(formula_sheet, column),
                    }
                )
    used_range = None
    if bounds is not None:
        min_col, min_row, max_col, max_row = bounds
        used_range = (
            f"{get_column_letter(min_col)}{min_row}:"
            f"{get_column_letter(max_col)}{max_row}"
        )
    return {
        "name": formula_sheet.title,
        "index": formula_sheet.parent.sheetnames.index(formula_sheet.title),
        "state": formula_sheet.sheet_state,
        "used_range": used_range,
        "merged_ranges": [str(item) for item in formula_sheet.merged_cells.ranges],
        "cells": cells,
        "tables": detect_tables(formula_sheet, bounds),
    }


def extract_workbook(source: str | Path) -> dict[str, Any]:
    """Extract ``source`` without calculating or rendering formulas."""

    path = Path(source).expanduser().resolve()
    if path.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise ValueError(f"unsupported workbook type: {path.suffix or '<none>'}")
    keep_vba = path.suffix.casefold() == ".xlsm"
    formulas = openpyxl.load_workbook(path, data_only=False, keep_vba=keep_vba)
    cached = openpyxl.load_workbook(path, data_only=True, keep_vba=keep_vba)
    try:
        sheets = [
            _extract_sheet(formulas[name], cached[name]) for name in formulas.sheetnames
        ]
        epoch_year = formulas.epoch.year
        return {
            "schema_version": SCHEMA_VERSION,
            "parser": {
                "name": PARSER_NAME,
                "library": "openpyxl",
                "library_version": openpyxl.__version__,
            },
            "source": {
                "filename": path.name,
                "path": str(path),
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                "size_bytes": path.stat().st_size,
            },
            "workbook": {
                "date_system": "1904" if epoch_year == 1904 else "1900",
                "sheet_count": len(sheets),
                "sheets": sheets,
            },
        }
    finally:
        formulas.close()
        cached.close()
