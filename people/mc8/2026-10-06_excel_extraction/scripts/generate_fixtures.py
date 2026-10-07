#!/usr/bin/env python3
"""Generate small, shareable workbooks for extraction and evaluation tests."""

from __future__ import annotations

import json
import zipfile
from datetime import date, datetime
from pathlib import Path
from tempfile import NamedTemporaryFile
from xml.etree import ElementTree

from openpyxl import Workbook
from openpyxl.styles import Font


ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "fixtures"
MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
FIXED_TIMESTAMP = (2026, 10, 6, 0, 0, 0)


def make_workbook() -> Workbook:
    """Create a workbook with deterministic document metadata."""

    workbook = Workbook()
    fixed_datetime = datetime(*FIXED_TIMESTAMP)
    workbook.properties.created = fixed_datetime
    workbook.properties.modified = fixed_datetime
    workbook.properties.creator = "MSCF document parsing capstone"
    workbook.properties.lastModifiedBy = "MSCF document parsing capstone"
    return workbook


def inject_numeric_formula_caches(path: Path, caches: dict[int, dict[str, float]]) -> None:
    """Inject cached results into generated OOXML without evaluating formulas.

    ``openpyxl`` intentionally does not calculate formulas and clears cached
    results when writing. This helper is fixture-generation code, not part of
    the parser. Sheet numbers follow workbook order.
    """

    ElementTree.register_namespace("", MAIN_NS)
    with NamedTemporaryFile(suffix=".xlsx", delete=False, dir=path.parent) as temporary:
        temporary_path = Path(temporary.name)
    with zipfile.ZipFile(path, "r") as source:
        with zipfile.ZipFile(temporary_path, "w") as destination:
            for member in source.infolist():
                content = source.read(member.filename)
                if member.filename.startswith("xl/worksheets/sheet"):
                    sheet_number = int(Path(member.filename).stem.removeprefix("sheet"))
                    sheet_caches = caches.get(sheet_number, {})
                    if sheet_caches:
                        root = ElementTree.fromstring(content)
                        for cell in root.findall(f".//{{{MAIN_NS}}}c"):
                            coordinate = cell.attrib.get("r")
                            if coordinate not in sheet_caches:
                                continue
                            formula = cell.find(f"{{{MAIN_NS}}}f")
                            if formula is None:
                                raise ValueError(f"{coordinate} is not a formula cell")
                            value = cell.find(f"{{{MAIN_NS}}}v")
                            if value is None:
                                value = ElementTree.SubElement(cell, f"{{{MAIN_NS}}}v")
                            value.text = str(sheet_caches[coordinate])
                            cell.attrib.pop("t", None)
                        content = ElementTree.tostring(
                            root, encoding="utf-8", xml_declaration=False
                        )
                member.date_time = FIXED_TIMESTAMP
                destination.writestr(member, content)
    temporary_path.replace(path)


def _style_header(sheet: object, row: int, last_column: int) -> None:
    for column in range(1, last_column + 1):
        sheet.cell(row, column).font = Font(bold=True)


def build_edge_cases(path: Path) -> None:
    workbook = make_workbook()
    sheet = workbook.active
    sheet.title = "Financial Statement"

    sheet.merge_cells("A1:D1")
    sheet["A1"] = "Illustrative Income Statement"
    sheet["A1"].font = Font(bold=True)
    sheet.append([])
    sheet.append(["Line item", 2025, 2024, "Review note"])
    _style_header(sheet, 3, 4)
    sheet.append(["Revenue", 1200, 1000, "audited"])
    sheet.append(["Expenses", 750, 625, "hidden row"])
    sheet.append(["Operating profit", "=B4-B5", "=C4-C5", "#DIV/0!"])
    sheet["D6"].data_type = "e"
    sheet.row_dimensions[5].hidden = True
    sheet.column_dimensions["D"].hidden = True
    sheet["B4"].number_format = "$#,##0"
    sheet["B5"].number_format = "$#,##0"
    sheet["B6"].number_format = "$#,##0"
    sheet.merge_cells("A8:D8")
    sheet["A8"] = "Note: B6 has a cache; C6 deliberately does not."
    sheet["A8"].font = Font(italic=True)

    sheet.merge_cells("A11:C11")
    sheet["A11"] = "Key Ratios"
    sheet["A11"].font = Font(bold=True)
    sheet.append(["Ratio", 2025, 2024])
    _style_header(sheet, 12, 3)
    sheet.append(["Operating margin", 0.375, 0.375])
    sheet.append(["Expense ratio", 0.625, 0.625])
    sheet.merge_cells("A16:C16")
    sheet["A16"] = "Source: generated test fixture"

    hidden = workbook.create_sheet("Hidden Inputs")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "Assumption"
    hidden["B1"] = 0.05
    hidden["A2"] = "Use forecast"
    hidden["B2"] = True
    hidden["A3"] = "As of"
    hidden["B3"] = date(2026, 10, 6)
    hidden["B3"].number_format = "yyyy-mm-dd"

    workbook.save(path)
    inject_numeric_formula_caches(path, {1: {"B6": 450}})


def build_gold_statement(path: Path) -> None:
    workbook = make_workbook()
    balance = workbook.active
    balance.title = "Balance Sheet"
    balance.merge_cells("A1:C1")
    balance["A1"] = "Example Limited — Balance Sheet"
    balance["A1"].font = Font(bold=True)
    balance.append([])
    balance.append(["Line item", 2025, 2024])
    _style_header(balance, 3, 3)
    balance.append(["Cash and cash equivalents", 1_250_000, 1_110_000])
    balance.append(["Property, plant and equipment", 2_500_000, 2_300_000])
    balance.append(["Total assets", "=SUM(B4:B5)", "=SUM(C4:C5)"])
    for row in range(4, 7):
        for column in (2, 3):
            balance.cell(row, column).number_format = "£#,##0"
    balance.merge_cells("A8:C8")
    balance["A8"] = "Note: amounts shown in GBP"

    income = workbook.create_sheet("Profit and Loss")
    income.merge_cells("A1:C1")
    income["A1"] = "Example Limited — Profit and Loss"
    income["A1"].font = Font(bold=True)
    income.append([])
    income.append(["Line item", 2025, 2024])
    _style_header(income, 3, 3)
    income.append(["Revenue", 4_800_000, 4_200_000])
    income.append(["Cost of sales", -2_900_000, -2_600_000])
    income.append(["Gross profit", "=SUM(B4:B5)", "=SUM(C4:C5)"])
    for row in range(4, 7):
        for column in (2, 3):
            income.cell(row, column).number_format = "£#,##0;[Red]-£#,##0"
    income.merge_cells("A8:C8")
    income["A8"] = "Source: generated and hand-labeled for this benchmark"

    workbook.save(path)
    inject_numeric_formula_caches(
        path,
        {
            1: {"B6": 3_750_000, "C6": 3_410_000},
            2: {"B6": 1_900_000, "C6": 1_600_000},
        },
    )


def write_gold_annotations(path: Path) -> None:
    payload = {
        "workbook": "gold_financial_statement.xlsx",
        "description": "Hand-labeled line-item survival checks; not a financial ontology.",
        "records": [
            {
                "sheet": "Balance Sheet",
                "label_coordinate": "A4",
                "label": "Cash and cash equivalents",
                "values": [
                    {"coordinate": "B4", "field": "value", "expected": 1250000},
                    {"coordinate": "C4", "field": "value", "expected": 1110000}
                ]
            },
            {
                "sheet": "Balance Sheet",
                "label_coordinate": "A6",
                "label": "Total assets",
                "values": [
                    {"coordinate": "B6", "field": "cached_value", "expected": 3750000},
                    {"coordinate": "C6", "field": "cached_value", "expected": 3410000}
                ]
            },
            {
                "sheet": "Profit and Loss",
                "label_coordinate": "A4",
                "label": "Revenue",
                "values": [
                    {"coordinate": "B4", "field": "value", "expected": 4800000},
                    {"coordinate": "C4", "field": "value", "expected": 4200000}
                ]
            },
            {
                "sheet": "Profit and Loss",
                "label_coordinate": "A5",
                "label": "Cost of sales",
                "values": [
                    {"coordinate": "B5", "field": "value", "expected": -2900000},
                    {"coordinate": "C5", "field": "value", "expected": -2600000}
                ]
            },
            {
                "sheet": "Profit and Loss",
                "label_coordinate": "A6",
                "label": "Gross profit",
                "values": [
                    {"coordinate": "B6", "field": "cached_value", "expected": 1900000},
                    {"coordinate": "C6", "field": "cached_value", "expected": 1600000}
                ]
            }
        ]
    }
    path.write_text(json.dumps(payload, indent=2), encoding="utf-8")


def main() -> None:
    FIXTURES.mkdir(parents=True, exist_ok=True)
    build_edge_cases(FIXTURES / "edge_cases.xlsx")
    build_gold_statement(FIXTURES / "gold_financial_statement.xlsx")
    write_gold_annotations(FIXTURES / "gold_annotations.json")
    print(f"Generated fixtures in {FIXTURES}")


if __name__ == "__main__":
    main()
