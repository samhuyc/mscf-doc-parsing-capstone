#!/usr/bin/env python3
"""Build deterministic Excel cases from the sponsor pilot and TAT-QA tables."""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from openpyxl.styles import Font


ROOT = Path(__file__).resolve().parents[1]
PEOPLE_ROOT = ROOT.parents[1]
DEFAULT_SPONSOR_SOURCE = (
    PEOPLE_ROOT / "samhu" / "2026-10-05_pilot_benchmark" / "reference_tables"
)
DEFAULT_OUTPUT = ROOT / "benchmark_cases"
sys.path.insert(0, str(ROOT))

from scripts.generate_fixtures import (  # noqa: E402
    inject_numeric_formula_caches,
    make_workbook,
)


TATQA_COMMIT = "644770eb2a66dddc24b92303bd2acbad84cd2b9f"
TATQA_SHA256 = "8da095a819af6db3c14877c6df2d4d29960e41d1a63dd1fa853507bd2a616af5"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


@dataclass
class HtmlCell:
    text: str
    tag: str
    rowspan: int
    colspan: int


class ReferenceTableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.rows: list[list[HtmlCell]] = []
        self.current_row: list[HtmlCell] | None = None
        self.current_cell: dict[str, Any] | None = None

    def handle_starttag(
        self, tag: str, attrs: list[tuple[str, str | None]]
    ) -> None:
        if tag == "tr":
            self.current_row = []
        elif tag in {"th", "td"}:
            attributes = dict(attrs)
            self.current_cell = {
                "tag": tag,
                "rowspan": int(attributes.get("rowspan") or 1),
                "colspan": int(attributes.get("colspan") or 1),
                "parts": [],
            }

    def handle_data(self, data: str) -> None:
        if self.current_cell is not None:
            self.current_cell["parts"].append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag in {"th", "td"} and self.current_cell is not None:
            if self.current_row is None:
                raise ValueError("table cell found outside a row")
            text = " ".join("".join(self.current_cell["parts"]).split())
            self.current_row.append(
                HtmlCell(
                    text=text,
                    tag=self.current_cell["tag"],
                    rowspan=self.current_cell["rowspan"],
                    colspan=self.current_cell["colspan"],
                )
            )
            self.current_cell = None
        elif tag == "tr" and self.current_row is not None:
            self.rows.append(self.current_row)
            self.current_row = None


def sponsor_case(path: Path, index: int) -> dict[str, Any]:
    parser = ReferenceTableParser()
    parser.feed(path.read_text(encoding="utf-8"))
    occupied: dict[tuple[int, int], str | None] = {}
    merges: list[dict[str, int]] = []
    header_rows = 0
    found_data = False
    for row_index, row in enumerate(parser.rows, start=1):
        if not found_data and not any(cell.tag == "td" for cell in row):
            header_rows += 1
        else:
            found_data = True
        column_index = 1
        for cell in row:
            while (row_index, column_index) in occupied:
                column_index += 1
            occupied[(row_index, column_index)] = cell.text
            if cell.rowspan > 1 or cell.colspan > 1:
                merges.append(
                    {
                        "min_row": row_index,
                        "min_column": column_index,
                        "max_row": row_index + cell.rowspan - 1,
                        "max_column": column_index + cell.colspan - 1,
                    }
                )
            for target_row in range(row_index, row_index + cell.rowspan):
                for target_column in range(column_index, column_index + cell.colspan):
                    occupied.setdefault((target_row, target_column), None)
            column_index += cell.colspan

    row_count = max(row for row, _ in occupied)
    column_count = max(column for _, column in occupied)
    values = [
        [occupied.get((row, column)) for column in range(1, column_count + 1)]
        for row in range(1, row_count + 1)
    ]
    return {
        "sheet": f"T{index:03d}",
        "source_id": path.name,
        "title": path.stem,
        "rows": row_count,
        "columns": column_count,
        "header_rows": header_rows,
        "values": values,
        "merges": merges,
    }


def build_sponsor_gold(source: Path) -> dict[str, Any]:
    paths = sorted(source.glob("*.html"))
    if not paths:
        raise FileNotFoundError(f"no sponsor reference tables found in {source}")
    return {
        "schema_version": "1.0.0",
        "benchmark": {
            "id": "sponsor-pilot-reference-tables",
            "description": "Eleven preliminary reference tables from the October 5 sponsor pilot.",
            "source_kind": "HTML tables derived from assistant-authored draft labels",
            "warning": "The references are not independently human-verified gold.",
            "conversion": "Each table becomes one worksheet; HTML row/column spans become merged cells.",
            "source_files": [
                {"name": path.name, "sha256": sha256(path)} for path in paths
            ],
        },
        "cases": [sponsor_case(path, index) for index, path in enumerate(paths, start=1)],
    }


def build_tatqa_gold(source: Path) -> dict[str, Any]:
    actual_sha256 = sha256(source)
    if actual_sha256 != TATQA_SHA256:
        raise ValueError(
            f"TAT-QA SHA-256 is {actual_sha256}; expected pinned {TATQA_SHA256}"
        )
    dataset = json.loads(source.read_text(encoding="utf-8"))
    cases = []
    for index, item in enumerate(dataset, start=1):
        source_rows = item["table"]["table"]
        columns = max(len(row) for row in source_rows)
        values = [
            [row[column] if column < len(row) else "" for column in range(columns)]
            for row in source_rows
        ]
        cases.append(
            {
                "sheet": f"T{index:03d}",
                "source_id": item["table"]["uid"],
                "title": item["table"]["uid"],
                "rows": len(values),
                "columns": columns,
                "header_rows": 1,
                "values": values,
                "merges": [],
            }
        )
    return {
        "schema_version": "1.0.0",
        "benchmark": {
            "id": "tatqa-development-tables",
            "description": "All 278 financial tables from the official English TAT-QA development split.",
            "source_kind": "Published TAT-QA JSON table structures",
            "license": "CC BY 4.0",
            "upstream_commit": TATQA_COMMIT,
            "source_sha256": actual_sha256,
            "conversion": "Each table becomes one worksheet; the context UID becomes a merged title row.",
        },
        "cases": cases,
    }


def write_workbook(gold: dict[str, Any], destination: Path) -> None:
    workbook = make_workbook()
    workbook.remove(workbook.active)
    for case in gold["cases"]:
        sheet = workbook.create_sheet(case["sheet"])
        if case["columns"] > 1:
            sheet.merge_cells(
                start_row=1,
                start_column=1,
                end_row=1,
                end_column=case["columns"],
            )
        sheet.cell(1, 1).value = case["title"]
        sheet.cell(1, 1).font = Font(bold=True)
        for row_index, row in enumerate(case["values"], start=1):
            for column_index, value in enumerate(row, start=1):
                if value is not None:
                    cell = sheet.cell(row_index + 2, column_index)
                    cell.value = value
                    if row_index <= case["header_rows"]:
                        cell.font = Font(bold=True)
        for merged in case["merges"]:
            sheet.merge_cells(
                start_row=merged["min_row"] + 2,
                start_column=merged["min_column"],
                end_row=merged["max_row"] + 2,
                end_column=merged["max_column"],
            )
    workbook.save(destination)
    inject_numeric_formula_caches(destination, {})


def write_case_set(gold: dict[str, Any], stem: str, output: Path) -> None:
    gold_path = output / f"{stem}_gold.json"
    workbook_path = output / f"{stem}.xlsx"
    gold_path.write_text(json.dumps(gold, ensure_ascii=False, indent=2), encoding="utf-8")
    write_workbook(gold, workbook_path)
    print(f"Wrote {len(gold['cases'])} cases: {workbook_path}")
    print(f"Wrote gold data: {gold_path}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--sponsor-source", type=Path, default=DEFAULT_SPONSOR_SOURCE)
    parser.add_argument(
        "--tatqa-source",
        type=Path,
        required=True,
        help="Pinned tatqa_dataset_dev.json with SHA-256 verification",
    )
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    args.output_dir.mkdir(parents=True, exist_ok=True)
    write_case_set(
        build_sponsor_gold(args.sponsor_source), "sponsor_pilot_tables", args.output_dir
    )
    write_case_set(
        build_tatqa_gold(args.tatqa_source), "tatqa_dev_tables", args.output_dir
    )


if __name__ == "__main__":
    main()
