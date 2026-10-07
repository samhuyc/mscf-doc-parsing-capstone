#!/usr/bin/env python3
"""Evaluate the Excel parser on the committed sponsor and TAT-QA cases."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CASES = ROOT / "benchmark_cases"
DEFAULT_OUTPUT = ROOT / "results" / "benchmark_cases.json"
sys.path.insert(0, str(ROOT))

from excel_parser import extract_workbook  # noqa: E402


def expected_range(case: dict[str, Any]) -> str:
    return f"A3:{get_column_letter(case['columns'])}{case['rows'] + 2}"


def evaluate_case_set(workbook_path: Path, gold_path: Path) -> dict[str, Any]:
    gold = json.loads(gold_path.read_text(encoding="utf-8"))
    payload = extract_workbook(workbook_path)
    sheets = {sheet["name"]: sheet for sheet in payload["workbook"]["sheets"]}
    totals = {
        "tables": len(gold["cases"]),
        "cells": 0,
        "coordinates_correct": 0,
        "nonblank_values": 0,
        "nonblank_values_correct": 0,
        "explicit_blank_values": 0,
        "explicit_blanks_preserved_as_empty_string": 0,
        "explicit_blanks_normalized_to_null": 0,
        "merged_cells": 0,
        "merged_cells_correct": 0,
        "single_table_detected": 0,
        "table_ranges_correct": 0,
        "titles_correct": 0,
        "header_classifications_correct": 0,
    }
    failures: list[dict[str, Any]] = []

    for case in gold["cases"]:
        sheet = sheets[case["sheet"]]
        by_coordinate = {cell["coordinate"]: cell for cell in sheet["cells"]}
        value_errors: list[dict[str, Any]] = []
        coordinate_errors: list[str] = []
        for row_index, row in enumerate(case["values"], start=1):
            for column_index, expected in enumerate(row, start=1):
                totals["cells"] += 1
                coordinate = f"{get_column_letter(column_index)}{row_index + 2}"
                actual_cell = by_coordinate.get(coordinate)
                if actual_cell is None:
                    coordinate_errors.append(coordinate)
                    continue
                totals["coordinates_correct"] += 1
                actual = actual_cell["value"]
                if expected == "":
                    totals["explicit_blank_values"] += 1
                    totals["explicit_blanks_preserved_as_empty_string"] += int(actual == "")
                    totals["explicit_blanks_normalized_to_null"] += int(actual is None)
                elif expected is not None:
                    totals["nonblank_values"] += 1
                    matched = actual == expected
                    totals["nonblank_values_correct"] += int(matched)
                    if not matched and len(value_errors) < 10:
                        value_errors.append(
                            {
                                "coordinate": coordinate,
                                "expected": expected,
                                "actual": actual,
                            }
                        )

        merge_errors: list[str] = []
        for merged in case["merges"]:
            anchor = (
                f"{get_column_letter(merged['min_column'])}{merged['min_row'] + 2}"
            )
            range_ref = (
                f"{anchor}:"
                f"{get_column_letter(merged['max_column'])}{merged['max_row'] + 2}"
            )
            for row in range(merged["min_row"], merged["max_row"] + 1):
                for column in range(
                    merged["min_column"], merged["max_column"] + 1
                ):
                    totals["merged_cells"] += 1
                    coordinate = f"{get_column_letter(column)}{row + 2}"
                    actual = by_coordinate[coordinate]["merged"]
                    expected_role = "anchor" if coordinate == anchor else "child"
                    correct = bool(
                        actual
                        and actual["range"] == range_ref
                        and actual["anchor"] == anchor
                        and actual["role"] == expected_role
                    )
                    totals["merged_cells_correct"] += int(correct)
                    if not correct:
                        merge_errors.append(coordinate)

        one_table = len(sheet["tables"]) == 1
        totals["single_table_detected"] += int(one_table)
        table = sheet["tables"][0] if one_table else None
        range_correct = bool(table and table["range"] == expected_range(case))
        title_correct = bool(table and table["title"] == case["title"])
        expected_headers = list(range(3, 3 + case["header_rows"]))
        actual_headers = table["row_roles"]["header"] if table else []
        headers_correct = actual_headers == expected_headers
        totals["table_ranges_correct"] += int(range_correct)
        totals["titles_correct"] += int(title_correct)
        totals["header_classifications_correct"] += int(headers_correct)

        if (
            value_errors
            or coordinate_errors
            or merge_errors
            or not one_table
            or not range_correct
            or not title_correct
            or not headers_correct
        ):
            failures.append(
                {
                    "source_id": case["source_id"],
                    "dimensions": f"{case['rows']}x{case['columns']}",
                    "value_errors": value_errors,
                    "coordinate_errors": coordinate_errors[:10],
                    "merge_errors": merge_errors[:10],
                    "expected_range": expected_range(case),
                    "detected_ranges": [item["range"] for item in sheet["tables"]],
                    "expected_header_rows": expected_headers,
                    "detected_header_rows": actual_headers,
                    "title_correct": title_correct,
                }
            )

    denominator = totals["nonblank_values"] + totals["explicit_blank_values"]
    strict_matches = (
        totals["nonblank_values_correct"]
        + totals["explicit_blanks_preserved_as_empty_string"]
    )
    return {
        "benchmark": gold["benchmark"],
        "workbook": workbook_path.name,
        "totals": totals,
        "rates": {
            "coordinate_accuracy": totals["coordinates_correct"] / totals["cells"],
            "nonblank_value_accuracy": (
                totals["nonblank_values_correct"] / totals["nonblank_values"]
            ),
            "strict_value_accuracy_including_empty_strings": strict_matches / denominator,
            "merged_cell_accuracy": (
                totals["merged_cells_correct"] / totals["merged_cells"]
                if totals["merged_cells"]
                else None
            ),
            "single_table_detection": totals["single_table_detected"] / totals["tables"],
            "exact_table_range": totals["table_ranges_correct"] / totals["tables"],
            "title_accuracy": totals["titles_correct"] / totals["tables"],
            "header_classification_accuracy": (
                totals["header_classifications_correct"] / totals["tables"]
            ),
        },
        "failure_count": len(failures),
        "failures": failures,
    }


def evaluate_all(case_dir: Path) -> dict[str, Any]:
    pairs = [
        ("sponsor_pilot_tables.xlsx", "sponsor_pilot_tables_gold.json"),
        ("tatqa_dev_tables.xlsx", "tatqa_dev_tables_gold.json"),
    ]
    return {
        "method": (
            "Reference tables were converted into XLSX worksheets, extracted to canonical "
            "JSON, and compared with their source cell matrices and layout annotations."
        ),
        "important_limit": (
            "These cases test realistic financial-table content in controlled XLSX "
            "containers; neither benchmark supplies native source Excel workbooks."
        ),
        "results": [
            evaluate_case_set(case_dir / workbook, case_dir / gold)
            for workbook, gold in pairs
        ],
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--case-dir", type=Path, default=DEFAULT_CASES)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--print-json", action="store_true")
    args = parser.parse_args()

    result = evaluate_all(args.case_dir)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    if args.print_json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        for benchmark in result["results"]:
            totals = benchmark["totals"]
            print(
                f"{benchmark['benchmark']['id']}: "
                f"nonblank values {totals['nonblank_values_correct']}/"
                f"{totals['nonblank_values']}; exact table ranges "
                f"{totals['table_ranges_correct']}/{totals['tables']}; "
                f"header rows {totals['header_classifications_correct']}/"
                f"{totals['tables']}"
            )
        print(f"Wrote {args.output}")

    for benchmark in result["results"]:
        totals = benchmark["totals"]
        if totals["coordinates_correct"] != totals["cells"]:
            raise SystemExit(1)
        if totals["nonblank_values_correct"] != totals["nonblank_values"]:
            raise SystemExit(1)
        if totals["merged_cells_correct"] != totals["merged_cells"]:
            raise SystemExit(1)


if __name__ == "__main__":
    main()
